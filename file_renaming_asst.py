import argparse
import base64
import io
import json
import logging
import os
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional, Sequence

import openai
import pandas as pd
import PyPDF2
import pytesseract
from docx import Document
from openpyxl import load_workbook
from pdf2image import convert_from_path
from PIL import Image


SCRIPT_DIR = Path(__file__).resolve().parent
CREDS_PATH = SCRIPT_DIR / "credentials.json"
DEFAULT_MODEL = "gpt-5.2"
DEFAULT_VISION_MODEL = "gpt-4o"
FILENAME_MAX_LEN = 50
DATE_PREFIX_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2} - ")
ISO_DATE_PATTERN = re.compile(r"(\d{4})[-/](\d{2})[-/](\d{2})")
SHORT_DATE_PATTERN = re.compile(
    r"(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})"
)  # for MM/DD/YYYY or similar
MONTH_NAME_PATTERN = re.compile(
    r"\b("
    r"January|February|March|April|May|June|July|August|September|October|November|December|"
    r"Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec"
    r")\s+(\d{1,2})(?:st|nd|rd|th)?[\s,]+(\d{2,4})",
    re.IGNORECASE,
)


asst_instructions = """You help users rename files by generating a concise and descriptive new name based on the file text content.

IMPORTANT NAMING RULES:
1. Always prepend the filename with a date in YYYY-MM-DD format representing the earliest date found in the document content (dates, filing dates, effective dates, etc.), followed by " - ". Use only dates you can clearly see in the document text.
2. Use spaces between words instead of underscores or hyphens.
3. NEVER include business names, company names, or organization names in the filename as these are redundant.
4. Keep the total filename length under 50 characters including spaces and the file extension.
5. Focus on the document TYPE and PURPOSE rather than who created it.
Respond with JSON only: {"new_name": "<proposed filename including extension>"}.
"""


@dataclass
class ExtractionOptions:
    percent: int = 100  # percentage of the leading content to use
    use_pdf_images: bool = True
    use_vision: bool = True
    max_preview_chars: Optional[int] = None


@dataclass
class Pricing:
    input_per_million: float
    output_per_million: float


MODEL_PRICING: dict[str, Pricing] = {
    # Override via CLI flags if these differ from your account pricing.
    "gpt-5.2": Pricing(input_per_million=5.0, output_per_million=15.0),
    "gpt-4o": Pricing(input_per_million=5.0, output_per_million=15.0),
}


class FileRenamerAssistant:
    def __init__(
        self,
        client: openai.OpenAI,
        *,
        options: ExtractionOptions,
        verbose: bool = False,
        dry_run: bool = False,
        model: str = DEFAULT_MODEL,
        vision_model: str = DEFAULT_VISION_MODEL,
        pricing: Optional[Pricing] = None,
    ):
        self.client = client
        self.options = options
        self.dry_run = dry_run
        self.model = model
        self.vision_model = vision_model
        self.pricing = pricing
        self.logger = logging.getLogger("file_renamer")
        self.logger.setLevel(logging.DEBUG if verbose else logging.INFO)

    def rename_directory(self, dir_path: Path) -> None:
        if not dir_path.is_dir():
            raise ValueError(f"Not a directory: {dir_path}")

        files = sorted(
            [p for p in dir_path.iterdir() if p.is_file() and not p.name.startswith(".")]
        )
        if not files:
            self.logger.info("No files found to rename in %s", dir_path)
            return

        self.logger.info(
            "Processing %d files in %s%s",
            len(files),
            dir_path,
            " (dry run)" if self.dry_run else "",
        )

        for file_path in files:
            try:
                self._rename_single(file_path)
            except Exception as exc:  # noqa: BLE001
                self.logger.error("Skipping %s: %s", file_path.name, exc)

    def _rename_single(self, file_path: Path) -> None:
        original_name = file_path.name
        content = self._extract_text_from_file(file_path)
        if self.options.max_preview_chars and len(content) > self.options.max_preview_chars:
            content = content[: self.options.max_preview_chars] + "...[truncated]"

        proposed_name = self._propose_name(original_name, content)
        validated_name = self._validate_name(
            proposed_name, original_name, content, file_path.suffix
        )

        if validated_name == original_name:
            self.logger.info("No change for %s (name already compliant)", original_name)
            return

        target_path = file_path.with_name(validated_name)
        if target_path.exists():
            raise FileExistsError(f"Target file already exists: {target_path.name}")

        if self.dry_run:
            self.logger.info("[dry run] %s -> %s", original_name, validated_name)
        else:
            file_path.rename(target_path)
            self.logger.info("Renamed: %s -> %s", original_name, validated_name)

    def _propose_name(self, original_name: str, content: str) -> str:
        prompt = f"""
Current filename: "{original_name}"
Document content (truncated): {content}

Return JSON only with the proposed new filename (including extension) under the key "new_name". Follow all naming rules strictly.
"""
        response = self.client.responses.create(
            model=self.model,
            input=[
                {"role": "system", "content": asst_instructions},
                {"role": "user", "content": prompt},
            ],
        )

        self._log_usage_cost(response)
        output_text = getattr(response, "output_text", "") or self._first_text_output(
            response
        )
        if not output_text:
            raise RuntimeError("Empty response from model.")

        return self._extract_name_from_response(output_text.strip())

    def _extract_name_from_response(self, output_text: str) -> str:
        try:
            data = json.loads(output_text)
            if isinstance(data, dict) and "new_name" in data:
                return str(data["new_name"]).strip()
        except json.JSONDecodeError:
            pass

        match = re.search(r'"new_name"\s*:\s*"([^"]+)"', output_text)
        if match:
            return match.group(1).strip()

        first_line = output_text.splitlines()[0].strip()
        if first_line:
            return first_line
        raise ValueError("Could not parse new filename from model output.")

    def _validate_name(
        self, name: str, original_name: str, content: str, original_ext: str
    ) -> str:
        cleaned = name.strip().replace("_", " ").replace("-", " ")
        base, ext = os.path.splitext(cleaned)
        if not ext:
            ext = original_ext
        elif ext.lower() != original_ext.lower():
            base = cleaned[: -len(ext)]
            ext = original_ext

        base = self._strip_existing_date(base)
        base_clean = base.strip() or Path(original_name).stem
        base_clean = self._strip_leading_numbers(base_clean)
        base_clean = re.sub(r"\s+", " ", base_clean).strip()
        date_prefix = (
            self._earliest_date(content)
            or self._date_in_string(base_clean)
            or self._date_in_string(original_name)
            or "Date Unknown"
        )

        final_name = self._assemble_with_length(date_prefix, base_clean, ext)
        return final_name

    def _strip_existing_date(self, name: str) -> str:
        return DATE_PREFIX_PATTERN.sub("", name).strip()

    def _strip_leading_numbers(self, text: str) -> str:
        return re.sub(r"^[0-9\s.-]+", "", text).strip()

    def _assemble_with_length(self, date_prefix: str, base: str, ext: str) -> str:
        separator = " - "
        budget = FILENAME_MAX_LEN - len(ext) - len(date_prefix) - len(separator)
        if budget < 1:
            raise ValueError("Filename budget too small for required components.")
        if len(base) > budget:
            base = base[:budget].rsplit(" ", 1)[0] or base[:budget]
        return f"{date_prefix}{separator}{base}{ext}"

    def _earliest_date(self, text: str) -> Optional[str]:
        candidates: list[datetime] = []

        for match in ISO_DATE_PATTERN.finditer(text):
            try:
                candidates.append(
                    datetime(year=int(match.group(1)), month=int(match.group(2)), day=int(match.group(3)))
                )
            except ValueError:
                continue

        for match in SHORT_DATE_PATTERN.finditer(text):
            month, day, year_raw = match.groups()
            year = int(year_raw)
            if year < 100:
                year += 2000 if year < 50 else 1900
            try:
                candidates.append(datetime(year=year, month=int(month), day=int(day)))
            except ValueError:
                continue

        month_map = {
            "january": 1,
            "february": 2,
            "march": 3,
            "april": 4,
            "may": 5,
            "june": 6,
            "july": 7,
            "august": 8,
            "september": 9,
            "october": 10,
            "november": 11,
            "december": 12,
            "jan": 1,
            "feb": 2,
            "mar": 3,
            "apr": 4,
            "jun": 6,
            "jul": 7,
            "aug": 8,
            "sep": 9,
            "sept": 9,
            "oct": 10,
            "nov": 11,
            "dec": 12,
        }
        for match in MONTH_NAME_PATTERN.finditer(text):
            month_name, day, year_raw = match.groups()
            month_num = month_map.get(month_name.lower())
            if not month_num:
                continue
            year = int(year_raw)
            if year < 100:
                year += 2000 if year < 50 else 1900
            try:
                candidates.append(datetime(year=year, month=month_num, day=int(day)))
            except ValueError:
                continue

        spaced = re.compile(r"(\d{4})[ .](\d{2})[ .](\d{2})")
        for match in spaced.finditer(text):
            try:
                candidates.append(
                    datetime(year=int(match.group(1)), month=int(match.group(2)), day=int(match.group(3)))
                )
            except ValueError:
                continue

        dotted = re.compile(r"(\d{4})[.](\d{2})[.](\d{2})")
        for match in dotted.finditer(text):
            try:
                candidates.append(
                    datetime(year=int(match.group(1)), month=int(match.group(2)), day=int(match.group(3)))
                )
            except ValueError:
                continue

        if not candidates:
            return None
        return min(candidates).strftime("%Y-%m-%d")

    def _date_in_string(self, text: str) -> Optional[str]:
        date = self._earliest_date(text)
        return date

    def _extract_text_from_file(self, file_path: Path) -> str:
        ext = file_path.suffix.lower()
        if ext == ".txt":
            return self._extract_text_from_txt(file_path)
        if ext == ".pdf":
            return self._extract_text_from_pdf(file_path)
        if ext == ".docx":
            return self._extract_text_from_docx(file_path)
        if ext == ".xlsx":
            return self._extract_text_from_xlsx(file_path)
        if ext == ".csv":
            return self._extract_text_from_csv(file_path)
        if ext in {".jpg", ".jpeg", ".png"}:
            return self._extract_text_from_image(file_path)
        raise ValueError(f"Unsupported file format: {ext}")

    def _slice_count(self, total: int) -> int:
        count = max(1, int(total * self.options.percent / 100))
        return count

    def _extract_text_from_txt(self, file_path: Path) -> str:
        with file_path.open(encoding="utf-8") as file:
            lines = file.readlines()
        return "".join(lines[: self._slice_count(len(lines))])

    def _extract_text_from_pdf(self, file_path: Path) -> str:
        try:
            with file_path.open("rb") as file:
                reader = PyPDF2.PdfReader(file)
                total_pages = len(reader.pages)
                page_count = min(self._slice_count(total_pages), total_pages)
                text = []
                for idx in range(page_count):
                    page_text = reader.pages[idx].extract_text() or ""
                    if page_text.strip():
                        text.append(page_text)

            extracted = "\n".join(text).strip()
            if extracted or not self.options.use_pdf_images:
                return extracted or "[No extractable text found]"

            images = convert_from_path(file_path, first_page=1, last_page=min(5, total_pages))
            vision_texts = []
            for img in images[:3]:
                img_bytes = io.BytesIO()
                img.save(img_bytes, format="JPEG")
                img_bytes = img_bytes.getvalue()
                page_text = self._extract_text_with_vision(img_bytes)
                if page_text.strip():
                    vision_texts.append(page_text)
            return "\n\n".join(vision_texts) if vision_texts else "[Scanned PDF - no readable text found]"
        except Exception as exc:  # noqa: BLE001
            self.logger.error("Error reading PDF %s: %s", file_path.name, exc)
            return "[Error reading PDF file]"

    def _extract_text_from_docx(self, file_path: Path) -> str:
        doc = Document(file_path)
        paragraphs = doc.paragraphs
        count = self._slice_count(len(paragraphs))
        return "\n".join(para.text for para in paragraphs[:count])

    def _extract_text_from_xlsx(self, file_path: Path) -> str:
        wb = load_workbook(file_path, read_only=True)
        sheet = wb.active
        rows = list(sheet.rows)
        count = self._slice_count(len(rows))
        return "\n".join(
            str(cell.value) for row in rows[:count] for cell in row if cell.value is not None
        )

    def _extract_text_from_csv(self, file_path: Path) -> str:
        df = pd.read_csv(file_path)
        count = self._slice_count(len(df))
        return df.head(count).to_string(index=False)

    def _extract_text_from_image(self, file_path: Path) -> str:
        try:
            if self.options.use_vision:
                llm_text = self._extract_text_with_vision(file_path)
                if llm_text and llm_text.strip():
                    return llm_text
            image = Image.open(file_path)
            return pytesseract.image_to_string(image)
        except Exception as exc:  # noqa: BLE001
            self.logger.error("Error reading image %s: %s", file_path.name, exc)
            return "[Error reading image file]"

    def _extract_text_with_vision(self, image_path_or_bytes: Path | bytes) -> str:
        try:
            if isinstance(image_path_or_bytes, bytes):
                image_data = base64.b64encode(image_path_or_bytes).decode("utf-8")
            else:
                with open(image_path_or_bytes, "rb") as f:
                    image_data = base64.b64encode(f.read()).decode("utf-8")

            response = self.client.responses.create(
                model=self.vision_model,
                input=[
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "input_text",
                                "text": "Extract all text content from this image. Return only the text content without commentary.",
                            },
                            {
                                "type": "input_image",
                                "image_url": f"data:image/jpeg;base64,{image_data}",
                            },
                        ],
                    }
                ],
                max_output_tokens=4000,
            )
            raw_text = getattr(response, "output_text", "") or self._first_text_output(response)
            return self._clean_vision_text(raw_text)
        except Exception as exc:  # noqa: BLE001
            self.logger.error("Vision extraction failed: %s", exc)
            return ""

    def _first_text_output(self, response) -> str:
        try:
            for item in response.output[0].content:
                if getattr(item, "type", "") == "output_text":
                    return item.text
        except Exception:
            return ""
        return ""

    def _clean_vision_text(self, text: str) -> str:
        """Remove obvious noise/base64-like blobs from vision output."""
        if not text:
            return ""
        stripped = text.strip()
        whitespace_ratio = sum(1 for c in stripped if c.isspace()) / max(len(stripped), 1)
        base64ish = re.fullmatch(r"[A-Za-z0-9+/=\s]+", stripped) is not None
        long_and_dense = len(stripped) > 500 and whitespace_ratio < 0.05
        if base64ish and long_and_dense:
            return ""
        return stripped

    def _log_usage_cost(self, response) -> None:
        usage = getattr(response, "usage", None)
        if not usage:
            return
        input_tokens = getattr(usage, "input_tokens", None) or usage.get("input_tokens", 0)
        output_tokens = getattr(usage, "output_tokens", None) or usage.get("output_tokens", 0)
        model_name = getattr(response, "model", None) or self.model

        cost_str = ""
        if self.pricing:
            input_cost = (input_tokens / 1_000_000) * self.pricing.input_per_million
            output_cost = (output_tokens / 1_000_000) * self.pricing.output_per_million
            total_cost = input_cost + output_cost
            cost_str = (
                f" | est cost ${total_cost:.6f} "
                f"(in ${input_cost:.6f}, out ${output_cost:.6f})"
            )

        self.logger.info(
            "Usage [%s]: input=%s tokens, output=%s tokens%s",
            model_name,
            input_tokens,
            output_tokens,
            cost_str,
        )


def load_creds() -> dict:
    if not CREDS_PATH.exists():
        raise FileNotFoundError(f"Missing credentials file: {CREDS_PATH}")
    with CREDS_PATH.open() as fh:
        return json.load(fh)


def build_client() -> openai.OpenAI:
    creds = load_creds()
    return openai.OpenAI(api_key=creds["openai_api_key"])


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Rename files in a directory based on content using OpenAI Responses API."
    )
    parser.add_argument("--files_rename", "-fr", type=Path, help="Directory containing files to rename")
    parser.add_argument("--dry_run", "-dr", action="store_true", help="Preview renames without writing")
    parser.add_argument("--verbose", "-v", action="store_true", help="Verbose logging")
    parser.add_argument(
        "--model",
        "-m",
        type=str,
        default=DEFAULT_MODEL,
        help="OpenAI model to use (default: gpt-5.2)",
    )
    parser.add_argument(
        "--vision_model",
        type=str,
        default=DEFAULT_VISION_MODEL,
        help="OpenAI vision-capable model for OCR fallback (default: gpt-4o)",
    )
    parser.add_argument(
        "--extraction_percent",
        "-p",
        type=int,
        default=100,
        help="Percent of file to sample from the start (1-100)",
    )
    parser.add_argument(
        "--max_preview_chars",
        "-mpc",
        type=int,
        default=None,
        help="Maximum characters of extracted text to send to the model (omit for no limit)",
    )
    parser.add_argument(
        "--disable_pdf_images",
        action="store_true",
        help="Skip image-based extraction for PDFs",
    )
    parser.add_argument(
        "--disable_vision",
        action="store_true",
        help="Skip vision model extraction for images",
    )
    parser.add_argument(
        "--price_in",
        type=float,
        help="Override input cost per 1M tokens for the chosen model",
    )
    parser.add_argument(
        "--price_out",
        type=float,
        help="Override output cost per 1M tokens for the chosen model",
    )
    parser.add_argument(
        "--openai_log_level",
        type=str,
        choices=["debug", "info", "warning", "error", "critical", "none"],
        default="none",
        help="OpenAI client log level (default: none to suppress HTTP debug).",
    )
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> None:
    args = parse_args(argv)
    if not args.files_rename:
        raise SystemExit("Please provide --files_rename <directory>.")

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s %(message)s",
    )

    openai.log = None if args.openai_log_level == "none" else args.openai_log_level

    options = ExtractionOptions(
        percent=max(1, min(args.extraction_percent, 100)),
        use_pdf_images=not args.disable_pdf_images,
        use_vision=not args.disable_vision,
        max_preview_chars=args.max_preview_chars,
    )

    pricing = MODEL_PRICING.get(args.model)
    if args.price_in is not None or args.price_out is not None:
        pricing = Pricing(
            input_per_million=args.price_in or (pricing.input_per_million if pricing else 0.0),
            output_per_million=args.price_out or (pricing.output_per_million if pricing else 0.0),
        )

    client = build_client()
    renamer = FileRenamerAssistant(
        client,
        options=options,
        verbose=args.verbose,
        dry_run=args.dry_run,
        model=args.model,
        vision_model=args.vision_model,
        pricing=pricing,
    )
    renamer.rename_directory(args.files_rename)


if __name__ == "__main__":
    main()
